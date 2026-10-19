import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

FINO = Side(style='thin', color='000000')
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
AZUL = PatternFill('solid', fgColor='4472C4')
ROT = PatternFill('solid', fgColor='B4C6E7')
LARANJA = PatternFill('solid', fgColor='FFC000')
VERDE = PatternFill('solid', fgColor='00B050')
CLARO = PatternFill('solid', fgColor='DDEBF7')
CINZA = PatternFill('solid', fgColor='D6DCE4')
CINZA2 = PatternFill('solid', fgColor='D9D9D9')
VERMELHO = PatternFill('solid', fgColor='FF0000')
CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)

# (chave, rótulo, formato)
LINHAS = (('ordem', 'Nº Ordem', None), ('planejado', 'PLANEJADO', '#,##0'), ('produzido', 'PRODUZIDO', '#,##0'),
          ('pacotes', 'TOTAL DE PACOTES', '#,##0'), ('mr0036', 'MR0036 CONSUMIDOS', '#,##0.00'),
          ('mescla', 'MESCLA FARINHA', None), ('absorcao', 'ABSORÇÃO DE ÁGUA', None),
          ('mr0001', 'MR0001 CONSUMIDOS', '#,##0.00'), ('mr0036g', 'MR0036 GERADOS', '#,##0.00'),
          ('mr0027', 'MR0027 CONSUMIDOS', '#,##0.00'), ('gelo', 'PACOTES DE GELO', '#,##0'))


def _fmt(n):
    return ('%g' % float(n)).replace('.', ',')


def _estiliza(c, fill=None, fonte=None, fmt=None, alinh=CENTRO):
    c.border = BORDA; c.alignment = alinh
    if fill: c.fill = fill
    c.font = fonte or Font(bold=True, size=12)
    if fmt: c.number_format = fmt


def _bloco(ws, r0, par, d, primeiro):
    """Um bloco do report (até 2 produtos lado a lado, como no modelo)."""
    ws.merge_cells(start_row=r0, end_row=r0, start_column=1, end_column=5)
    c = ws.cell(row=r0, column=1, value='REPORT DE PRODUÇÃO DIÁRIA JBT')
    for col in range(1, 6): _estiliza(ws.cell(row=r0, column=col), AZUL, Font(bold=True, size=18, color='FFFFFF'))
    ws.row_dimensions[r0].height = 36
    data = '/'.join(reversed(str(d.get('data', '')).split('-')))
    # turno / data / observação
    _estiliza(ws.cell(row=r0 + 1, column=1, value='TURNO'), ROT)
    ws.merge_cells(start_row=r0 + 1, end_row=r0 + 1, start_column=2, end_column=4)
    for col in (2, 3, 4): _estiliza(ws.cell(row=r0 + 1, column=col), VERDE)
    ws.cell(row=r0 + 1, column=2).value = str(d.get('turno', '')).upper()
    _estiliza(ws.cell(row=r0 + 1, column=5, value='OBSERVAÇÃO'), CINZA2)
    _estiliza(ws.cell(row=r0 + 2, column=1, value='DATA'), ROT)
    ws.merge_cells(start_row=r0 + 2, end_row=r0 + 2, start_column=2, end_column=4)
    for col in (2, 3, 4): _estiliza(ws.cell(row=r0 + 2, column=col), CLARO)
    ws.cell(row=r0 + 2, column=2).value = data
    _estiliza(ws.cell(row=r0 + 2, column=5, value='total de gelo em estoque'), ROT, Font(size=12))
    # produto
    r = r0 + 3
    _estiliza(ws.cell(row=r, column=1, value='PRODUTO'), ROT)
    for col in (2, 4):
        p = par[0 if col == 2 else 1] if len(par) > (0 if col == 2 else 1) else None
        _estiliza(ws.cell(row=r, column=col, value=(p.get('codigo') or p.get('produto')) if p else None), LARANJA)
    _estiliza(ws.cell(row=r, column=3), CINZA)
    est = d.get('estoque_gelo')
    _estiliza(ws.cell(row=r, column=5, value=(f'{_fmt(est)} pallet' if primeiro and est not in (None, '') else None)), ROT, Font(size=16))
    for i, (k, rot, fmt) in enumerate(LINHAS, 1):
        rr = r + i
        _estiliza(ws.cell(row=rr, column=1, value=rot), ROT)
        for col, idx in ((2, 0), (4, 1)):
            v = par[idx].get(k) if len(par) > idx else None
            _estiliza(ws.cell(row=rr, column=col, value=v if v not in (None, '') else None), LARANJA, fmt=fmt)
        _estiliza(ws.cell(row=rr, column=3), CINZA)
        _estiliza(ws.cell(row=rr, column=5), ROT, Font(size=12))
    # observações (coluna E)
    if primeiro:
        ws.merge_cells(start_row=r + 2, end_row=r + len(LINHAS), start_column=5, end_column=5)
        c = ws.cell(row=r + 2, column=5, value=d.get('obs') or None)
        c.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
    # absenteísmo
    ra = r + len(LINHAS) + 1
    _estiliza(ws.cell(row=ra, column=1, value='ABSENTEÍSMO'), ROT)
    ws.merge_cells(start_row=ra, end_row=ra, start_column=2, end_column=4)
    for col in (2, 3, 4): _estiliza(ws.cell(row=ra, column=col), VERMELHO, Font(bold=True, size=12, color='FFFFFF'))
    _estiliza(ws.cell(row=ra, column=5), ROT)
    if primeiro:
        aus = d.get('ausentes', []); tc = int(d.get('total_colab') or 0)
        if tc or aus:
            pct = _fmt(round(len(aus) / tc * 100, 1)) if tc else '0'
            txt = f'{len(aus)} de {tc} ({pct}%)'
            if aus: txt += ': ' + '; '.join(f"{a.get('nome')} ({a.get('motivo')})" for a in aus)
            ws.cell(row=ra, column=2).value = txt
            ws.row_dimensions[ra].height = 30 if len(txt) < 90 else 60
    return ra


def gerar_report(d):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Report de produção'
    for col, w in zip('ABCDE', (28, 30, 6, 30, 32)): ws.column_dimensions[col].width = w
    prods = d.get('produtos', []) or [{}]
    pares = [prods[i:i + 2] for i in range(0, len(prods), 2)]
    r0, quebras = 1, []
    for n, par in enumerate(pares):
        fim = _bloco(ws, r0, par, d, n == 0)
        quebras.append(fim)
        r0 = fim + 2
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_area = f'A1:E{quebras[-1]}'
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()
