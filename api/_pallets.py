import copy
import io
import os

import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.pagebreak import Break

AQUI = os.path.dirname(os.path.abspath(__file__))
# (linha inicial, quantidade de linhas) dos dois blocos da ficha: lotes 1-34 / 35-68 e 69-99 / 100-130
BLOCOS = ((12, 34), (48, 31))
ESQ = dict(turno='B', hora='C', lote='D', qtd='E', letra='F', resp='H', obs='I')
DIR = dict(turno='J', hora='K', lote='L', qtd='M', letra='N', resp='P', obs='Q')
CAPACIDADE = 130
ESCALA = 50  # % de impressão: cada metade da ficha (linhas 1-45 e 46-81) cabe inteira em uma folha A4 paisagem


def _imprimir(ws):
    """Configura a impressão: 2 folhas por produto, quebrando entre os dois blocos da ficha.
    Sem 'ajustar à página' (o Excel ignora quebras manuais nesse modo e cortava a tabela no meio)."""
    ws.sheet_properties.pageSetUpPr = None
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9
    ws.page_setup.scale = ESCALA
    ws.page_setup.fitToWidth = None
    ws.page_setup.fitToHeight = None
    ws.print_area = 'A1:R81'
    ws.print_title_rows = '2:10'
    ws.print_options.horizontalCentered = True
    ws.row_breaks.brk = []
    ws.row_breaks.append(Break(id=45))


def _pos(i):
    """Posição (colunas, linha) do i-ésimo pallet da folha: esquerda e direita de cada bloco."""
    for ini, n in BLOCOS:
        if i < n: return ESQ, ini + i
        i -= n
        if i < n: return DIR, ini + i
        i -= n
    return None, None


def _w(ws, ref, valor, tam=11, negrito=True, alinh='center'):
    c = ws[ref]; c.value = valor
    c.font = Font(name='Calibri', size=tam, bold=negrito)
    c.alignment = Alignment(horizontal=alinh, vertical='center', wrap_text=True)


def _folha(ws, produto, data, itens):
    ws['C10'].value = produto
    ws['P10'].value = data
    # linhas sem pallet voltam ao visual padrão do formulário em branco
    for ini, n in BLOCOS:
        for r in range(ini, ini + n):
            ws[f'B{r}'].value = '    1          2         3'
            ws[f'J{r}'].value = ' 1             2            3'
    for i, p in enumerate(itens[:CAPACIDADE]):
        cols, r = _pos(i)
        _w(ws, f"{cols['turno']}{r}", str(p.get('turno', ''))[:1] + 'º')
        _w(ws, f"{cols['hora']}{r}", p.get('horario'))
        _w(ws, f"{cols['lote']}{r}", p.get('lote'), 14)
        _w(ws, f"{cols['qtd']}{r}", p.get('quantidade'))
        _w(ws, f"{cols['letra']}{r}", p.get('letra'), 14)
        _w(ws, f"{cols['resp']}{r}", p.get('responsavel'), 10, False)
        obs = (p.get('tipo') or '').upper()
        if p.get('obs'): obs = f"{obs} - {p['obs']}"
        _w(ws, f"{cols['obs']}{r}", obs, 9, False)


def gerar_pallets(d):
    """Ficha FM-000082 preenchida com os pallets do dia; uma folha por produto."""
    wb = openpyxl.load_workbook(os.path.join(AQUI, '_modelo_pallets.xlsx'))
    base = wb['JBT']
    img = base._images[0] if base._images else None
    raw, anchor = (img._data(), copy.deepcopy(img.anchor)) if img else (None, None)
    if raw:
        base._images = []
        i0 = Image(io.BytesIO(raw)); i0.anchor = copy.deepcopy(anchor); base.add_image(i0)
    data = '/'.join(reversed(str(d.get('data', '')).split('-')))
    grupos = {}
    for p in sorted(d.get('pallets', []), key=lambda x: x.get('lote', 0)):
        grupos.setdefault((p.get('codigo') or '', p.get('produto') or ''), []).append(p)
    if not grupos:
        grupos = {('', ''): []}
    folhas = []
    for k, (chave, itens) in enumerate(grupos.items()):
        ws = base if k == 0 else wb.copy_worksheet(base)
        if k and raw:
            im = Image(io.BytesIO(raw)); im.anchor = copy.deepcopy(anchor); ws.add_image(im)
        _imprimir(ws)
        nome = (chave[0] or chave[1] or 'JBT')[:28]
        folhas.append((ws, nome, chave, itens))
    # primeiro copia as folhas em branco (acima), depois preenche e nomeia
    usados = set()
    for ws, nome, (cod, prod), itens in folhas:
        t, n = nome, 2
        while t in usados: t = f'{nome[:26]} {n}'; n += 1
        usados.add(t); ws.title = t
        _folha(ws, ' - '.join(x for x in (cod, prod) if x), data, itens)
    wb.active = wb.worksheets.index(folhas[0][0])
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()
