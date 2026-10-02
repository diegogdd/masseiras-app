import copy, io, os, re
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.pagebreak import Break

AQUI = os.path.dirname(os.path.abspath(__file__))
H = 50          # linhas de uma página do formulário
PRIM, ULT = 12, 48  # linhas de lotes na página (37 por página)

def _runs(bats):
    """Agrupa as bateladas por produto: novo grupo quando troca o produto ou o lote reinicia."""
    runs, ult = [], None
    for b in bats:
        if not runs or b['produto'] != runs[-1][0] or (ult is not None and b['lote'] < ult):
            runs.append((b['produto'], []))
        runs[-1][1].append(b)
        ult = b['lote']
    return [(p, sorted(l, key=lambda x: x['lote'])) for p, l in runs]

def _virgula(v):
    return ('%g' % float(v)).replace('.', ',') if v is not None else ''

def _clonar(ws, k, img_bytes, anchor):
    o = H * k
    for row in ws.iter_rows(min_row=1, max_row=H):
        for c in row:
            d = ws.cell(row=c.row + o, column=c.column, value=c.value)
            if c.has_style:
                d._style = copy.copy(c._style)
    for r in range(1, H + 1):
        ws.row_dimensions[r + o].height = ws.row_dimensions[r].height
    for m in list(ws.merged_cells.ranges):
        if m.max_row <= H:
            ws.merge_cells(start_row=m.min_row + o, end_row=m.max_row + o, start_column=m.min_col, end_column=m.max_col)
    if img_bytes:
        a = copy.deepcopy(anchor)
        a._from.row += o
        if hasattr(a, 'to'):
            a.to.row += o
        im = Image(io.BytesIO(img_bytes)); im.anchor = a
        ws.add_image(im)

def _marca(ws, ref, txt, tam=11):
    c = ws[ref]
    c.value = txt
    c.font = Font(name='Calibri', size=tam, bold=True)
    c.alignment = Alignment(horizontal='center', vertical='center')

def gerar(dados):
    wb = openpyxl.load_workbook(os.path.join(AQUI, '_modelo.xlsx'))
    ws = wb.worksheets[0]
    img0 = ws._images[0]
    raw, anchor = img0._data(), copy.deepcopy(img0.anchor)
    ws._images = []
    im0 = Image(io.BytesIO(raw)); im0.anchor = copy.deepcopy(anchor); ws.add_image(im0)
    pags = []  # (produto, lotes)
    for prod, lotes in _runs(dados.get('bateladas', [])):
        for i in range(0, len(lotes), ULT - PRIM + 1):
            pags.append((prod, lotes[i:i + ULT - PRIM + 1]))
    if not pags:
        pags = [('', [])]
    for k in range(1, len(pags)):
        _clonar(ws, k, raw, anchor)
    data = '/'.join(reversed(str(dados.get('data', '')).split('-')))
    for k, (prod, lotes) in enumerate(pags):
        o = H * k
        _marca(ws, f'B{9+o}', data)
        ws[f'E{9+o}'].value = dados.get('turno', '')
        ws[f'I{9+o}'].value = prod
        ws[f'I{9+o}'].font = Font(name='Calibri', size=11, bold=True)
        for i, b in enumerate(lotes):
            r = PRIM + i + o
            _marca(ws, f'A{r}', b['lote'])
            _marca(ws, f'B{r}', b.get('hi') or ':')
            _marca(ws, f'C{r}', b.get('hf') or ':')
            _marca(ws, f'D{r}', 'X')
            _marca(ws, f'F{r}', 'X')
            l, c = ('(L)' if b.get('rep_linha') else ' L '), ('(C)' if b.get('rep_congelado') else ' C ')
            ws[f'E{r}'].value = f'  {l}      {c}'
            if b.get('rep_linha') or b.get('rep_congelado'):
                ws[f'E{r}'].font = Font(name='Calibri', size=11, bold=True)
            _marca(ws, f'{"G" if b.get("diosna") == 1 else "H"}{r}', 'X')
            t = _virgula(b.get('temp'))
            if t:
                ws[f'I{r}'].value = f'{t} °C'
                ws[f'I{r}'].alignment = Alignment(horizontal='center', vertical='center')
            ws[f'K{r}'].value = b.get('operador') or None
        if k < len(pags) - 1:
            ws.row_breaks.append(Break(id=H * (k + 1)))
    ws.print_area = f'A1:M{H*len(pags)}'
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.title = 'Parâmetros das Amassadeiras'
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()
