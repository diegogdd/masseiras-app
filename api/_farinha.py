import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

def gerar_farinha(d):
    """Planilha separada com moinho e lote de farinha usados em cada batelada."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Farinha por lote'
    data = '/'.join(reversed(str(d.get('data', '')).split('-')))
    ws['A1'] = 'Mescla de farinha por lote'; ws['A1'].font = Font(size=14, bold=True)
    ws['A2'] = 'Data:'; ws['B2'] = data; ws['D2'] = 'Turno:'; ws['E2'] = d.get('turno', '')
    for c in ('A2', 'D2', 'B2', 'E2'): ws[c].font = Font(bold=True)
    cab = ['Produto', 'Lote', 'Hora início', 'Moega 1 - Moinho', 'Moega 1 - Lote da farinha', 'Moega 2 - Moinho', 'Moega 2 - Lote da farinha']
    borda = Border(*(Side(style='thin'),) * 4)
    for i, t in enumerate(cab, 1):
        c = ws.cell(row=4, column=i, value=t)
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='D9D9D9')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = borda
    r = 5
    for b in d.get('bateladas', []):
        f = {x.get('moega'): x for x in (b.get('farinha') or [])}
        linha = [b.get('produto'), b.get('lote'), b.get('hi') or '', (f.get(1) or {}).get('moinho', '-'), (f.get(1) or {}).get('lote', '-'), (f.get(2) or {}).get('moinho', '-'), (f.get(2) or {}).get('lote', '-')]
        for i, v in enumerate(linha, 1):
            c = ws.cell(row=r, column=i, value=v); c.border = borda
            c.alignment = Alignment(horizontal='left' if i == 1 else 'center', vertical='center')
        r += 1
    for col, w in zip('ABCDEFG', (26, 8, 12, 20, 24, 20, 24)): ws.column_dimensions[col].width = w
    ws.row_dimensions[4].height = 32
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = '4:4'; ws.freeze_panes = 'A5'
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()
